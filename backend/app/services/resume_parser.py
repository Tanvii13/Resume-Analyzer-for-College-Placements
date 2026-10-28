"""
Resume parsing, validation, and error handling service for NLP Sprint 1.

Provides reusable, structured text extraction and validation for PDF and DOCX
resume documents, with clear error classification and metadata extraction.
"""

import io
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Optional, Tuple
from zipfile import BadZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pydantic import BaseModel
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError, PdfReadError

try:
    from docx.opc.exceptions import PackageNotFoundError
except ImportError:
    class PackageNotFoundError(Exception):
        """Fallback if PackageNotFoundError cannot be imported directly."""
        pass


MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class ResumeParsingErrorCode:
    """Stable error codes for structured error reporting."""
    EMPTY_FILE = "EMPTY_FILE"
    UNSUPPORTED_FORMAT = "UNSUPPORTED_FORMAT"
    FILE_TOO_LARGE = "FILE_TOO_LARGE"
    INVALID_FILE = "INVALID_FILE"
    CORRUPTED_FILE = "CORRUPTED_FILE"
    ENCRYPTED_FILE = "ENCRYPTED_FILE"
    NO_READABLE_TEXT = "NO_READABLE_TEXT"
    PARSER_ERROR = "PARSER_ERROR"


class ResumeParsingError(Exception):
    """Base exception for all resume parsing and validation errors."""

    def __init__(
        self,
        message: str,
        error_code: str = ResumeParsingErrorCode.PARSER_ERROR,
        status_code: int = 400,
    ):
        super().__init__(message)
        self.message = message
        self.error_code = error_code
        self.status_code = status_code

    def to_dict(self) -> dict:
        return {
            "detail": self.message,
            "error_code": self.error_code,
        }


class ResumeValidationError(ResumeParsingError):
    """Raised when an uploaded file fails pre-extraction validation."""
    pass


class EmptyFileError(ResumeValidationError):
    def __init__(self, message: str = "The uploaded file is empty. Please choose another file."):
        super().__init__(
            message=message,
            error_code=ResumeParsingErrorCode.EMPTY_FILE,
            status_code=400,
        )


class UnsupportedFormatError(ResumeValidationError):
    def __init__(self, message: str = "Unsupported file format. Please upload a PDF or DOCX file."):
        super().__init__(
            message=message,
            error_code=ResumeParsingErrorCode.UNSUPPORTED_FORMAT,
            status_code=415,
        )


class FileTooLargeError(ResumeValidationError):
    def __init__(self, message: str = "The resume exceeds the 5 MB upload limit."):
        super().__init__(
            message=message,
            error_code=ResumeParsingErrorCode.FILE_TOO_LARGE,
            status_code=413,
        )


class InvalidFileError(ResumeValidationError):
    def __init__(self, message: str = "The file is invalid."):
        super().__init__(
            message=message,
            error_code=ResumeParsingErrorCode.INVALID_FILE,
            status_code=400,
        )


class CorruptedFileError(ResumeValidationError):
    def __init__(self, message: str = "The document is corrupted or unreadable. Please upload another file."):
        super().__init__(
            message=message,
            error_code=ResumeParsingErrorCode.CORRUPTED_FILE,
            status_code=400,
        )


class EncryptedFileError(ResumeValidationError):
    def __init__(self, message: str = "Password-protected PDFs are not supported."):
        super().__init__(
            message=message,
            error_code=ResumeParsingErrorCode.ENCRYPTED_FILE,
            status_code=400,
        )


class ResumeExtractionError(ResumeParsingError):
    """Raised when text extraction fails during or after parsing."""
    pass


class NoReadableTextError(ResumeExtractionError):
    def __init__(
        self,
        message: str = (
            "No readable text could be extracted from the document. "
            "Scanned or image-only documents requiring OCR are not supported."
        ),
    ):
        super().__init__(
            message=message,
            error_code=ResumeParsingErrorCode.NO_READABLE_TEXT,
            status_code=400,
        )


class ExtractionFailureError(ResumeExtractionError):
    def __init__(self, message: str = "An unexpected error occurred while parsing the document."):
        super().__init__(
            message=message,
            error_code=ResumeParsingErrorCode.PARSER_ERROR,
            status_code=422,
        )


class ResumeParseResult(BaseModel):
    """Structured result returned by the resume parser interface."""

    filename: str
    file_type: str
    text: str
    page_count: Optional[int] = None
    char_count: int = 0
    word_count: int = 0

    @classmethod
    def create(
        cls,
        filename: str,
        file_type: str,
        text: str,
        page_count: Optional[int] = None,
    ) -> "ResumeParseResult":
        char_count = len(text)
        word_count = len(text.split())
        return cls(
            filename=filename,
            file_type=file_type,
            text=text,
            page_count=page_count,
            char_count=char_count,
            word_count=word_count,
        )


def normalize_whitespace(text: Optional[str]) -> str:
    """
    Normalize whitespace, line breaks, and unicode characters in extracted text.

    - Preserves meaningful line breaks and paragraph separation.
    - Normalizes unicode spaces and special characters.
    - Strips trailing spaces, zero-width spaces, and collapses excessive blank lines.
    """
    if not text:
        return ""

    # Normalize unicode to NFKC (converts non-breaking spaces, compatibility characters)
    normalized = unicodedata.normalize("NFKC", text)

    # Strip zero-width characters, soft hyphens, and byte-order marks
    normalized = re.sub(r"[\u200b-\u200d\ufeff\u00ad]", "", normalized)

    # Standardize line endings to \n
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")

    # Clean individual lines (collapse multiple spaces/tabs, trim ends)
    lines = [
        re.sub(r"[ \t\f\v]+", " ", line).strip()
        for line in normalized.split("\n")
    ]

    # Rejoin lines
    rejoined = "\n".join(lines)

    # Collapse 3 or more consecutive newlines into 2 (\n\n)
    rejoined = re.sub(r"\n{3,}", "\n\n", rejoined)

    return rejoined.strip()


def has_readable_text(text: Optional[str]) -> bool:
    """
    Determine if the extracted text contains meaningful readable content.
    Returns True if at least one alphanumeric character exists.
    """
    if not text:
        return False
    return any(char.isalnum() for char in text)


def validate_resume_metadata(
    filename: Optional[str],
    file_bytes: bytes,
) -> Tuple[str, str]:
    """
    Validate filename and basic file size/extension properties.

    Returns:
        Tuple of (clean_filename, normalized_extension).
    """
    if not filename or not str(filename).strip():
        raise InvalidFileError("Please select a resume file.")

    clean_filename = Path(str(filename)).name.strip()
    if not clean_filename or clean_filename in SUPPORTED_EXTENSIONS:
        raise InvalidFileError("Please select a resume file.")

    extension = Path(clean_filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(
            "Unsupported file format. Please upload a PDF or DOCX file."
        )

    if not file_bytes or len(file_bytes) == 0:
        raise EmptyFileError("The uploaded file is empty. Please choose another file.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise FileTooLargeError("The resume exceeds the 5 MB upload limit.")

    return clean_filename, extension


def validate_resume_structure(file_bytes: bytes, extension: str) -> None:
    """
    Validate the internal format and structure of the document according to its extension.
    Detects corruptions, mismatches, encrypted files, and non-DOCX ZIP files.
    """
    if extension == ".pdf":
        if file_bytes.startswith(b"PK"):
            raise InvalidFileError(
                "File extension is .pdf but content appears to be a ZIP or DOCX archive."
            )
        if not file_bytes.startswith(b"%PDF-"):
            raise InvalidFileError("The file does not appear to be a valid PDF.")

        try:
            reader = PdfReader(io.BytesIO(file_bytes), strict=False)
            if reader.is_encrypted:
                try:
                    res = reader.decrypt("")
                    if res == 0:
                        raise EncryptedFileError("Password-protected PDFs are not supported.")
                except EncryptedFileError:
                    raise
                except Exception:
                    raise EncryptedFileError("Password-protected PDFs are not supported.")
            _ = len(reader.pages)
        except (EncryptedFileError, InvalidFileError):
            raise
        except FileNotDecryptedError as exc:
            raise EncryptedFileError("Password-protected PDFs are not supported.") from exc
        except (PdfReadError, ValueError, KeyError, OSError, Exception) as exc:
            raise CorruptedFileError(
                "The document is corrupted or unreadable. Please upload another file."
            ) from exc

    elif extension == ".docx":
        if file_bytes.startswith(b"%PDF-"):
            raise InvalidFileError(
                "File extension is .docx but content appears to be a PDF document."
            )
        if file_bytes.startswith(b"\xd0\xcf\x11\xe0"):
            if b"EncryptedPackage" in file_bytes[:4096]:
                raise EncryptedFileError("Password-protected DOCX documents are not supported.")
            raise InvalidFileError("The file does not appear to be a valid DOCX document.")

        if not file_bytes.startswith(b"PK"):
            raise InvalidFileError("The file does not appear to be a valid DOCX document.")

        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                namelist = zf.namelist()
                if "word/document.xml" not in namelist:
                    raise InvalidFileError(
                        "The file does not appear to be a valid DOCX document."
                    )
            Document(io.BytesIO(file_bytes))
        except (InvalidFileError, CorruptedFileError, EncryptedFileError):
            raise
        except (BadZipFile, PackageNotFoundError, KeyError, ValueError, OSError, Exception) as exc:
            raise CorruptedFileError(
                "The document is corrupted or unreadable. Please upload another file."
            ) from exc
    else:
        raise UnsupportedFormatError(
            "Unsupported file format. Please upload a PDF or DOCX file."
        )


def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, int]:
    """
    Extract readable text from a PDF document.

    Returns:
        Tuple of (extracted_text, page_count).
    """
    try:
        reader = PdfReader(io.BytesIO(file_bytes), strict=False)
    except Exception as exc:
        raise CorruptedFileError(
            "The document is corrupted or unreadable. Please upload another file."
        ) from exc

    if reader.is_encrypted:
        try:
            res = reader.decrypt("")
            if res == 0:
                raise EncryptedFileError("Password-protected PDFs are not supported.")
        except EncryptedFileError:
            raise
        except Exception:
            raise EncryptedFileError("Password-protected PDFs are not supported.")

    try:
        page_count = len(reader.pages)
    except FileNotDecryptedError as exc:
        raise EncryptedFileError("Password-protected PDFs are not supported.") from exc
    except Exception as exc:
        raise CorruptedFileError(
            "The document is corrupted or unreadable. Please upload another file."
        ) from exc

    if page_count == 0:
        raise CorruptedFileError("The PDF document contains no pages.")

    pages_text: list[str] = []
    try:
        for idx, page in enumerate(reader.pages):
            try:
                raw_page_text = page.extract_text() or ""
            except Exception as exc:
                raise ExtractionFailureError(
                    f"Failed to extract text from PDF page {idx + 1}."
                ) from exc
            norm_page_text = normalize_whitespace(raw_page_text)
            if norm_page_text:
                pages_text.append(norm_page_text)
    except ResumeParsingError:
        raise
    except Exception as exc:
        raise ExtractionFailureError(
            "An unexpected error occurred while extracting text from the PDF document."
        ) from exc

    full_text = "\n\n".join(pages_text).strip()
    if not has_readable_text(full_text):
        raise NoReadableTextError(
            "No readable text could be extracted from the PDF. "
            "The document may be scanned, image-only, or blank. OCR is not supported."
        )

    return full_text, page_count


_W_TXBX_CONTENT = qn("w:txbxContent")
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _top_level_text_boxes(elem) -> list:
    """
    Return the text-box bodies (w:txbxContent) anchored inside a paragraph.

    Word stores each modern text box twice: once in mc:Choice and once as a
    legacy VML copy in mc:Fallback. Fallback copies are skipped to avoid
    duplicated text. Nested text boxes are skipped here because they are
    reached when their parent text box is processed.
    """
    boxes = []
    for box in elem.iter(_W_TXBX_CONTENT):
        ancestor = box.getparent()
        skip = False
        while ancestor is not None and ancestor is not elem:
            if ancestor.tag in (_MC_FALLBACK, _W_TXBX_CONTENT):
                skip = True
                break
            ancestor = ancestor.getparent()
        if not skip:
            boxes.append(box)
    return boxes


def _extract_from_docx_element(elem, doc: Document, parts: list[str]) -> None:
    """Helper to recursively extract text from DOCX elements preserving document order."""
    tag = elem.tag
    if tag.endswith("p"):
        para = Paragraph(elem, doc)
        p_text = normalize_whitespace(para.text)
        if p_text:
            parts.append(p_text)
        # Text boxes are common in resume templates (sidebars, contact blocks)
        # and are not included in Paragraph.text.
        for box in _top_level_text_boxes(elem):
            for inner in box.iterchildren():
                _extract_from_docx_element(inner, doc, parts)
    elif tag.endswith("tbl"):
        table = Table(elem, doc)
        for row in table.rows:
            seen_cells = set()
            row_cells = []
            for cell in row.cells:
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                cell_text = normalize_whitespace(cell.text)
                if cell_text:
                    row_cells.append(cell_text)
            if row_cells:
                parts.append(" | ".join(row_cells))
    elif tag.endswith("sdt"):
        for child in elem.iterchildren():
            if child.tag.endswith("sdtContent"):
                for inner in child.iterchildren():
                    _extract_from_docx_element(inner, doc, parts)
            else:
                _extract_from_docx_element(child, doc, parts)


def _extract_header_footer_text(doc: Document, kind: str) -> list[str]:
    """
    Extract text from every section's header or footer ("header" / "footer").

    Includes paragraphs, tables, and text boxes. Lines repeated across
    sections are kept once. Failures here never block body extraction.
    """
    lines: list[str] = []
    try:
        for section in doc.sections:
            part = getattr(section, kind)
            if part is None or part.is_linked_to_previous:
                continue
            found: list[str] = []
            for child in part._element.iterchildren():
                _extract_from_docx_element(child, doc, found)
            for line in found:
                if line not in lines:
                    lines.append(line)
    except Exception:
        pass
    return lines


def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, None]:
    """
    Extract readable text from a DOCX document.
    Preserves logical reading order of headers, paragraphs, content controls, and tables.

    Returns:
        Tuple of (extracted_text, None).
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
    except (BadZipFile, PackageNotFoundError, KeyError, ValueError, OSError, Exception) as exc:
        raise CorruptedFileError(
            "The document is corrupted or unreadable. Please upload another file."
        ) from exc

    parts: list[str] = []

    # Extract headers first (contact info and name often reside there)
    parts.extend(_extract_header_footer_text(doc, "header"))

    try:
        # Preserve document reading order by iterating through body elements
        for child in doc.element.body:
            _extract_from_docx_element(child, doc, parts)
    except Exception:
        # Fallback to standard paragraphs and tables iteration if element iteration fails
        try:
            parts = []
            for p in doc.paragraphs:
                p_text = normalize_whitespace(p.text)
                if p_text:
                    parts.append(p_text)
            for table in doc.tables:
                for row in table.rows:
                    seen_cells = set()
                    row_cells = []
                    for cell in row.cells:
                        if cell._tc in seen_cells:
                            continue
                        seen_cells.add(cell._tc)
                        cell_text = normalize_whitespace(cell.text)
                        if cell_text:
                            row_cells.append(cell_text)
                    if row_cells:
                        parts.append(" | ".join(row_cells))
        except Exception as exc:
            raise ExtractionFailureError(
                "An unexpected error occurred while parsing the DOCX document."
            ) from exc

    # Footers go last; some templates place links or contact details there
    for line in _extract_header_footer_text(doc, "footer"):
        if line not in parts:
            parts.append(line)

    full_text = "\n\n".join(parts).strip()
    if not has_readable_text(full_text):
        raise NoReadableTextError(
            "No readable text could be extracted from the DOCX document."
        )

    return full_text, None


def parse_resume(file_bytes: bytes, filename: str) -> ResumeParseResult:
    """
    Parse a resume document (PDF or DOCX) and return structured text and metadata.

    Args:
        file_bytes: Raw bytes of the uploaded resume document.
        filename: Original filename of the document.

    Returns:
        ResumeParseResult containing extracted text, page count, and character/word counts.

    Raises:
        ResumeValidationError: If the file fails metadata or structural validation.
        ResumeExtractionError: If text extraction fails or no readable text exists.
    """
    clean_filename, extension = validate_resume_metadata(filename, file_bytes)
    validate_resume_structure(file_bytes, extension)

    try:
        if extension == ".pdf":
            text, page_count = extract_text_from_pdf(file_bytes)
        elif extension == ".docx":
            text, page_count = extract_text_from_docx(file_bytes)
        else:
            raise UnsupportedFormatError(
                "Unsupported file format. Please upload a PDF or DOCX file."
            )
    except ResumeParsingError:
        raise
    except Exception as exc:
        raise ExtractionFailureError(
            "An unexpected error occurred while parsing the document."
        ) from exc

    return ResumeParseResult.create(
        filename=clean_filename,
        file_type=extension,
        text=text,
        page_count=page_count,
    )
