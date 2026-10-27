"""
Robustness tests for resume text extraction on real-world document quirks.

Covers:
- PDFs with a damaged cross-reference table (readable by every viewer)
- DOCX text boxes (sidebars/contact blocks in resume templates)
- DOCX text boxes not duplicated by Word's legacy mc:Fallback copy
- DOCX footers and header tables
"""

import io
import re
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml

backend_dir = str(Path(__file__).resolve().parents[1])
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from app.services.resume_parser import (  # noqa: E402
    CorruptedFileError,
    parse_resume,
)

SAMPLE_PDF = (
    Path(__file__).resolve().parents[2]
    / "frontend"
    / "sample_resumes"
    / "sample_resume.pdf"
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _text_box_paragraph(text: str) -> object:
    """
    Build a paragraph holding a text box the way Word saves one:
    a modern copy in mc:Choice and a legacy VML copy in mc:Fallback.
    """
    return parse_xml(
        f"""
        <w:p xmlns:w="{W_NS}"
             xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
             xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
             xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
             xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
             xmlns:v="urn:schemas-microsoft-com:vml">
          <w:r>
            <mc:AlternateContent>
              <mc:Choice Requires="wps">
                <w:drawing>
                  <wp:anchor>
                    <wp:extent cx="1" cy="1"/>
                    <wp:docPr id="1" name="TextBox"/>
                    <a:graphic>
                      <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                        <wps:wsp>
                          <wps:txbx>
                            <w:txbxContent>
                              <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
                            </w:txbxContent>
                          </wps:txbx>
                        </wps:wsp>
                      </a:graphicData>
                    </a:graphic>
                  </wp:anchor>
                </w:drawing>
              </mc:Choice>
              <mc:Fallback>
                <w:pict>
                  <v:shape>
                    <v:textbox>
                      <w:txbxContent>
                        <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
                      </w:txbxContent>
                    </v:textbox>
                  </v:shape>
                </w:pict>
              </mc:Fallback>
            </mc:AlternateContent>
          </w:r>
        </w:p>
        """
    )


def _docx_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def test_pdf_with_damaged_xref_is_still_parsed():
    """A wrong startxref offset is repaired by viewers and must not be rejected."""
    good = SAMPLE_PDF.read_bytes()
    damaged = re.sub(rb"startxref\s+\d+", b"startxref\n123", good)
    assert damaged != good

    result = parse_resume(damaged, "resume.pdf")

    assert "Alex Rivera" in result.text
    assert result.page_count == 1
    assert result.word_count == parse_resume(good, "resume.pdf").word_count


def test_truncated_pdf_is_still_rejected():
    """Lenient reading must not let genuinely broken files through."""
    good = SAMPLE_PDF.read_bytes()
    truncated = good[:200]

    with pytest.raises(CorruptedFileError):
        parse_resume(truncated, "resume.pdf")


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def test_docx_text_box_content_is_extracted():
    doc = Document()
    doc.add_paragraph("Alex Rivera")
    doc.element.body.insert(1, _text_box_paragraph("SKILLS Python Java SQL"))
    doc.add_paragraph("EXPERIENCE Intern at Acme Labs")

    result = parse_resume(_docx_bytes(doc), "resume.docx")

    assert "SKILLS Python Java SQL" in result.text


def test_docx_text_box_is_not_duplicated_from_fallback_copy():
    doc = Document()
    doc.add_paragraph("Alex Rivera")
    doc.element.body.insert(1, _text_box_paragraph("UniqueSidebarText"))

    result = parse_resume(_docx_bytes(doc), "resume.docx")

    assert result.text.count("UniqueSidebarText") == 1


def test_docx_text_box_keeps_reading_order():
    doc = Document()
    doc.add_paragraph("FIRST paragraph")
    doc.element.body.insert(1, _text_box_paragraph("SECOND in text box"))
    doc.add_paragraph("THIRD paragraph")

    text = parse_resume(_docx_bytes(doc), "resume.docx").text

    assert text.index("FIRST") < text.index("SECOND") < text.index("THIRD")


def test_docx_footer_is_extracted_after_body():
    doc = Document()
    doc.add_paragraph("Alex Rivera")
    doc.sections[0].footer.paragraphs[0].text = "github.com/alexrivera"

    text = parse_resume(_docx_bytes(doc), "resume.docx").text

    assert "github.com/alexrivera" in text
    assert text.index("Alex Rivera") < text.index("github.com/alexrivera")


def test_docx_header_table_is_extracted_before_body():
    doc = Document()
    doc.add_paragraph("EDUCATION B.Tech ICT")
    header = doc.sections[0].header
    table = header.add_table(rows=1, cols=2, width=doc.sections[0].page_width)
    table.cell(0, 0).text = "alex.rivera@dau.ac.in"
    table.cell(0, 1).text = "+91 98765 43210"

    text = parse_resume(_docx_bytes(doc), "resume.docx").text

    assert "alex.rivera@dau.ac.in | +91 98765 43210" in text
    assert text.index("alex.rivera") < text.index("EDUCATION")


def test_docx_repeated_header_across_sections_kept_once():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Alex Rivera - Resume"
    doc.add_paragraph("Page one content")
    new_section = doc.add_section()
    new_section.header.is_linked_to_previous = False
    new_section.header.paragraphs[0].text = "Alex Rivera - Resume"
    doc.add_paragraph("Page two content")

    text = parse_resume(_docx_bytes(doc), "resume.docx").text

    assert text.count("Alex Rivera - Resume") == 1
