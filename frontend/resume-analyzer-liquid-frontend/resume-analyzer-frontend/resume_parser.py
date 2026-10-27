"""
Resume parser — Sprint 1 (US-02 validation, US-03 parsing, US-04 invalid files)

Three steps, each a small function:
  1. validate_file()   -> is this a real, non-empty PDF/DOCX within the size limit?
  2. extract_text()    -> read the raw text out of the file
  3. split_sections()  -> organise the text into Education, Skills, Projects, ...

If anything goes wrong we raise ResumeParseError with an error CODE and a
plain-English MESSAGE, so the UI can tell the student exactly what to fix.
"""

import re
from pypdf import PdfReader
from docx import Document

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
MIN_TEXT_LENGTH = 50             # less text than this = probably a scanned image

# First bytes of each file type ("magic numbers"). These are checked because
# a user could simply rename virus.exe to resume.pdf.
FILE_SIGNATURES = {
    "pdf": b"%PDF-",
    "docx": b"PK\x03\x04",   # DOCX files are ZIP archives internally
}

# Words that usually appear as section headings in a resume
SECTION_HEADINGS = {
    "summary":        ["summary", "objective", "career objective", "profile", "about me", "professional summary"],
    "education":      ["education", "academic details", "academics", "academic qualifications", "qualifications"],
    "skills":         ["skills", "technical skills", "key skills", "core skills", "skills and tools", "technologies"],
    "experience":     ["experience", "work experience", "professional experience", "internship", "internships", "employment"],
    "projects":       ["projects", "academic projects", "personal projects", "key projects"],
    "certifications": ["certifications", "certificates", "courses", "licenses and certifications"],
    "achievements":   ["achievements", "awards", "honors", "honours", "accomplishments"],
    "activities":     ["extracurricular activities", "activities", "positions of responsibility", "leadership"],
}

SECTION_TITLES = {
    "summary": "Summary", "education": "Education", "skills": "Skills",
    "experience": "Experience", "projects": "Projects", "certifications": "Certifications",
    "achievements": "Achievements", "activities": "Activities",
}

# Sections a placement resume is expected to have
EXPECTED_SECTIONS = ["education", "skills", "projects", "experience"]

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?\d{5}[\s-]?\d{5}|\+?\d[\d\s-]{8,14}\d")
LINK_RE = re.compile(r"(?:https?://)?(?:www\.)?(?:linkedin\.com|github\.com)/[\w\-/]+", re.I)


class ResumeParseError(Exception):
    """Raised when a resume can't be accepted or read."""

    def __init__(self, code, message):
        super().__init__(message)
        self.code = code
        self.message = message


# ---------------------------------------------------------------------------
# Step 1: validation
# ---------------------------------------------------------------------------
def validate_file(filename, data):
    """Check name, size and real file type. Returns 'pdf' or 'docx'."""
    if not filename or "." not in filename:
        raise ResumeParseError("ERR_NO_FILE", "Please choose a file to upload.")

    ext = filename.rsplit(".", 1)[1].lower()
    if ext not in FILE_SIGNATURES:
        raise ResumeParseError("ERR_FILE_TYPE",
                               f"'.{ext}' files are not supported. Please upload a PDF or DOCX file.")
    if len(data) == 0:
        raise ResumeParseError("ERR_EMPTY_FILE", "This file is empty. Please upload your actual resume.")
    if len(data) > MAX_FILE_SIZE:
        raise ResumeParseError("ERR_FILE_TOO_LARGE", "This file is larger than 5 MB. Please compress it and try again.")
    if not data.startswith(FILE_SIGNATURES[ext]):
        raise ResumeParseError("ERR_FILE_CORRUPT",
                               f"This file is named .{ext} but its content is not a valid {ext.upper()}. "
                               "It may be corrupted or renamed.")
    return ext


# ---------------------------------------------------------------------------
# Step 2: text extraction
# ---------------------------------------------------------------------------
def extract_text(path, ext):
    """Return (text, page_count) from a PDF or DOCX file."""
    try:
        if ext == "pdf":
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ResumeParseError("ERR_PASSWORD_PROTECTED",
                                       "This PDF is password-protected. Please remove the password and upload again.")
            pages = [page.extract_text() or "" for page in reader.pages]
            text, page_count = "\n".join(pages), len(pages)
        else:
            doc = Document(path)
            lines = [p.text for p in doc.paragraphs]
            for table in doc.tables:              # many resumes use tables for layout
                for row in table.rows:
                    lines.extend(cell.text for cell in row.cells)
            text, page_count = "\n".join(lines), None
    except ResumeParseError:
        raise
    except Exception:  # any error from pypdf or python-docx means the file is damaged
        raise ResumeParseError("ERR_UNREADABLE",
                               "This file can't be read. It may be damaged. "
                               "Try exporting your resume again as PDF from Word or Google Docs.")

    if len(text.strip()) < MIN_TEXT_LENGTH:
        raise ResumeParseError("ERR_NO_TEXT",
                               "No readable text was found. This usually means the resume is a scanned image. "
                               "Please upload a text-based PDF or DOCX.")
    return text, page_count


# ---------------------------------------------------------------------------
# Step 3: organise into sections
# ---------------------------------------------------------------------------
def _heading_key(line):
    """If this line looks like a section heading, return its key (e.g. 'skills')."""
    cleaned = re.sub(r"[^a-z ]", "", line.lower()).strip()
    if not cleaned or len(cleaned) > 40:
        return None
    for key, words in SECTION_HEADINGS.items():
        if cleaned in words:
            return key
    return None


def split_sections(text):
    """Split resume text into a header (name/contact) and named sections."""
    header_lines, sections, current = [], {}, None

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        key = _heading_key(line)
        if key:
            current = key
            sections.setdefault(key, [])
        elif current:
            sections[current].append(line)
        else:
            header_lines.append(line)

    return header_lines, {k: "\n".join(v) for k, v in sections.items()}


def extract_contact(text, header_lines):
    """Find name, email, phone and profile links."""
    email = EMAIL_RE.search(text)
    phone = PHONE_RE.search(text)
    return {
        "name": header_lines[0] if header_lines else None,
        "email": email.group() if email else None,
        "phone": phone.group().strip() if phone else None,
        "links": sorted(set(LINK_RE.findall(text))),
    }


def parse_resume(path, ext):
    """Full pipeline used by app.py. Returns a dictionary the templates can show."""
    text, page_count = extract_text(path, ext)
    header_lines, sections = split_sections(text)

    return {
        "contact": extract_contact(text, header_lines),
        "sections": [
            {"key": k, "title": SECTION_TITLES[k], "content": v}
            for k, v in sections.items()
        ],
        "missing": [SECTION_TITLES[k] for k in EXPECTED_SECTIONS if k not in sections],
        "word_count": len(text.split()),
        "pages": page_count,
    }
